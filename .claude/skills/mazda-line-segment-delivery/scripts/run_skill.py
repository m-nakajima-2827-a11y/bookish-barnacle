"""mazda-line-segment-delivery skill を SKILL.md の仕様どおりに実行する（分析→分岐→Excel）。
使い方: python run_skill.py <UTF-8変換済みCSV> <元ファイル名> <出力xlsx> <出力json>
（CSVは事前に iconv -f CP932 -t UTF-8 で変換する。元ファイル名の regist_YYYYMMDD を基準日に使う）
"""
import sys, os, re, json
import pandas as pd
from pandas.tseries.offsets import DateOffset
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.dirname(__file__))
from texts import D

csv, srcname, out_xlsx, out_json = sys.argv[1:5]
m = re.search(r'regist_(\d{8})', srcname)
BASE = pd.Timestamp(m.group(1))
SRC = re.sub(r'^[0-9a-f]{8}-', '', os.path.basename(srcname))

# ---------- 1. 読み込みと前処理 ----------
df = pd.read_csv(csv, dtype=str, keep_default_na=False)
for c in df.columns:
    df[c] = df[c].astype(str).str.replace(r'^="?', '', regex=True).str.replace(r'"$', '', regex=True).str.strip()
df['店舗'] = df['担当店舗(管理者店舗)']
df['担当者'] = df['担当者'].replace({'サ-ビス': 'サービス', '': '(未設定)'})
reg, ins = pd.to_datetime(df['登録日']), pd.to_datetime(df['車検日'])
df['経過年数(現在)'] = ((BASE - reg).dt.days / 365.25).round(1)
nxt = ((ins - reg).dt.days / 365.25).round().astype(int)
df['次回車検の回数'] = pd.cut(nxt, [-1, 3, 5, 7, 9, 999], labels=['初回車検(3年)', '2回目(5年)', '3回目(7年)', '4回目(9年)', '10年超']).astype(str)
df['車検まで(月)'] = ((ins - BASE).dt.days / 30.44).round(1)
df['車検時期'] = pd.cut(df['車検まで(月)'], [-999, 2, 3, 6, 12, 9999], labels=['1) 2か月以内', '2) 2〜3か月', '3) 4〜6か月', '4) 7〜12か月', '5) 13か月以上']).astype(str)
df['複数台保有'] = df['電話番号'].map(df['電話番号'].value_counts()).gt(1).map({True: 'はい', False: ''})

# 1-3 車種の分類
CLS = [
 (['MAZDA2'], 'コンパクト', '現行', 'MAZDA2'), (['デミオ'], 'コンパクト', '旧名称', 'MAZDA2'),
 (['CX-3'], 'コンパクトSUV', '現行', 'CX-3／CX-30'),
 (['MAZDA3 FASTBACK', 'MAZDA3 SEDAN'], 'ミドル', '現行', 'MAZDA3／CX-30'),
 (['アクセラスポーツ', 'アクセラセダン', 'アクセラハイブリッド'], 'ミドル', '旧名称', 'MAZDA3／CX-30'),
 (['CX-30'], 'ミドルSUV', '現行', 'CX-30／CX-5'), (['CX-5'], 'ミドルSUV', '現行', 'CX-5／CX-60'),
 (['CX-60'], 'ラージSUV', '現行', 'CX-60／CX-80'), (['CX-80'], 'ラージSUV(3列)', '現行', 'CX-80'),
 (['CX-8'], 'ラージSUV(3列)', '販売終了', 'CX-80'), (['MX-30'], '電動SUV', '現行', 'MX-30／CX-60 PHEV'),
 (['ロードスター'], 'スポーツ', '現行', 'ロードスター'), (['RX-8', 'RX-7'], 'スポーツ', '販売終了', 'ロードスター'),
 (['アテンザワゴン', 'アテンザセダン', 'アテンザスポーツ', 'MAZDA6 SEDAN', 'MAZDA6 WAGON'], 'セダン・ワゴン', '販売終了', 'CX-5／CX-60'),
 (['プレマシー', 'ビアンテ', 'MPV'], 'ミニバン', '販売終了', 'CX-80／CX-5'),
 (['フレアワゴン'], '軽', '現行', 'フレアワゴン'), (['フレアクロスオーバー'], '軽', '現行', 'フレアクロスオーバー'),
 (['フレア'], '軽', '現行', 'フレア'), (['キャロル'], '軽', '現行', 'キャロル／フレア'), (['スクラムワゴン'], '軽', '現行', 'スクラムワゴン'),
 (['AZ-ワゴン'], '軽', '旧名称', 'フレア'),
]
MAP = {n: (k, z, p) for names, k, z, p in CLS for n in names}
def cls(v):
    if v in MAP: return MAP[v]
    return ('その他・他メーカー', 'その他', '用途を聞いてから提案') if v.startswith('その他') else ('未分類', '未分類', '')
df[['車格', '車種区分', '提案車種']] = pd.DataFrame(df['車種'].map(cls).tolist(), index=df.index)
unclassified = sorted(df.loc[df['車格'] == '未分類', '車種'].unique())

# 1-4 暫定価格帯セグメント
def seg(r):
    k, v = r['車格'], r['車種']
    if k in ('軽', 'コンパクト'): return '300万〜400万円'
    if k == 'コンパクトSUV': return '401万〜500万円'
    if k == 'ミドル' or v == 'CX-30': return '501万〜600万円'
    if v in ('CX-60', 'CX-8'): return '701万〜800万円'
    if v == 'CX-80': return '801万円以上'
    if v == 'CX-5' or k in ('スポーツ', 'セダン・ワゴン', 'ミニバン', '電動SUV'): return '601万〜700万円'
    return '未判定'
df['暫定価格帯セグメント(仮説)'] = df.apply(seg, axis=1)

# ---------- 3. 分岐 ----------
def prio(r):
    mm = r['車検まで(月)']
    if mm <= 6 and r['スカイプランの加入'] == 'はい': return 'A1 スカイプラン満了期（車検6か月以内）'
    if mm <= 6 and r['車種区分'] == '販売終了': return 'A2 販売終了車種（車検6か月以内）'
    if mm <= 6 and nxt[r.name] >= 7: return 'A3 7年以上（車検6か月以内）'
    if mm <= 6: return 'B 車検予約（車検6か月以内）'
    if mm <= 12: return 'C1 準備層（7〜12か月）'
    return 'C2 関係維持（13か月以上）'
df['優先度'] = df.apply(prio, axis=1)
STEPS = [('1回目', 6), ('2回目', 4), ('3回目', 3), ('4回目', 2)]
for n, mo in STEPS:
    d = ins - DateOffset(months=mo)
    df[f'{n}配信日'] = d.dt.strftime('%Y-%m-%d').where(d >= BASE, '経過')
def nextstep(r):
    for n, _ in STEPS:
        if r[f'{n}配信日'] != '経過': return n
    return '至急（車検2か月以内：4回目の内容を即時配信）'
df['次に送るステップ'] = df.apply(nextstep, axis=1)
df['至急'] = df['次に送るステップ'].str.startswith('至急')

# ---------- 集計 ----------
n = len(df)
PR = ['A1 スカイプラン満了期（車検6か月以内）', 'A2 販売終了車種（車検6か月以内）', 'A3 7年以上（車検6か月以内）', 'B 車検予約（車検6か月以内）', 'C1 準備層（7〜12か月）', 'C2 関係維持（13か月以上）']
STORES = list(df['店舗'].value_counts().index)
BUCK = ['1) 2か月以内', '2) 2〜3か月', '3) 4〜6か月', '4) 7〜12か月', '5) 13か月以上']
s6 = df['車検まで(月)'] <= 6
R = {}  # 数値の記録（計画書・改善点用）
R['n'] = n; R['persons'] = int(df['電話番号'].nunique()); R['multi'] = int((df['電話番号'].value_counts() > 1).sum())
R['stores'] = {s: int((df['店舗'] == s).sum()) for s in STORES}
for c, k in [('任意保険の加入', 'ins'), ('パックdeメンテの加入', 'pack'), ('スカイプランの加入', 'sky')]:
    R[k + '_rate'] = float((df[c] == 'はい').mean())
R['avg_age'] = float(df['経過年数(現在)'].mean()); R['s6'] = int(s6.sum()); R['s12'] = int((df['車検まで(月)'] <= 12).sum())

g = df.groupby('車種')
t2 = pd.DataFrame({'件数': g.size(), '平均経過年数': g['経過年数(現在)'].mean().round(1), '車検6か月以内': g['車検まで(月)'].apply(lambda s: int((s <= 6).sum())),
                   'スカイプラン加入': g['スカイプランの加入'].apply(lambda s: int((s == 'はい').sum())), 'パックdeメンテ加入率': g['パックdeメンテの加入'].apply(lambda s: round(float((s == 'はい').mean()), 2))})
t2 = t2.join(df.drop_duplicates('車種').set_index('車種')[['車格', '車種区分', '提案車種']]).sort_values('件数', ascending=False).reset_index()
g = df.groupby('車格')
t3 = pd.DataFrame({'件数': g.size(), '平均経過年数': g['経過年数(現在)'].mean().round(1), '車検6か月以内': g['車検まで(月)'].apply(lambda s: int((s <= 6).sum())),
                   '販売終了旧名称': g['車種区分'].apply(lambda s: int(s.isin(['販売終了', '旧名称']).sum()))}).sort_values('件数', ascending=False).reset_index()
t4 = pd.crosstab(df['店舗'], df['車検時期']).reindex(index=STORES, columns=BUCK, fill_value=0)
ORD = ['初回車検(3年)', '2回目(5年)', '3回目(7年)', '4回目(9年)', '10年超']
t5 = pd.crosstab(df['次回車検の回数'], df['店舗']).reindex(index=ORD, columns=STORES, fill_value=0)
cal = df[df['車検まで(月)'] <= 12].copy(); cal['満了月'] = pd.to_datetime(cal['車検日']).dt.strftime('%Y-%m')
t6 = pd.crosstab(cal['満了月'], cal['店舗']).reindex(columns=STORES, fill_value=0)
t7 = pd.crosstab(df['優先度'], df['店舗']).reindex(index=PR, columns=STORES, fill_value=0)
staff = df.groupby(['店舗', '担当者']).size().reset_index(name='件数').sort_values(['店舗', '件数'], ascending=[True, False])
staff['店舗'] = pd.Categorical(staff['店舗'], STORES); staff = staff.sort_values(['店舗', '件数'], ascending=[True, False])
AJ = ['任意保険の加入', 'パックdeメンテの加入', 'スカイプランの加入']
t9 = pd.concat({c: pd.crosstab(df['車検時期'], df[c]).reindex(index=BUCK, columns=['', 'いいえ', 'はい'], fill_value=0) for c in AJ}, axis=1)
SEGS = ['300万〜400万円', '401万〜500万円', '501万〜600万円', '601万〜700万円', '701万〜800万円', '801万円以上', '未判定']
t10 = pd.crosstab(df['暫定価格帯セグメント(仮説)'], df['車検時期']).reindex(index=SEGS, columns=BUCK, fill_value=0)
seg_models = {}
for sname, gg in df.groupby('暫定価格帯セグメント(仮説)'):
    vc = gg['車種'].value_counts()
    seg_models[sname] = (f'その他・他メーカー車（「その他(…)」で登録の車種 {len(vc)}種類）' if sname == '未判定'
                         else '、'.join(f'{k}（{v}）' for k, v in vc.items()))
t11 = pd.crosstab(df['誕生月'], df['店舗']).reindex(index=[f'{i}月' for i in range(1, 13)], columns=STORES, fill_value=0)
t12 = pd.crosstab(df['友だち登録年'] if '友だち登録年' in df else pd.to_datetime(df['お友だち登録日']).dt.year, df['店舗']).reindex(columns=STORES, fill_value=0)
other = df[df['車種'].str.startswith('その他')]['車種']
mazda_in_other = sorted({v for v in other if re.search(r'デミオ|DEMIO|アクセラ|アテンザ|ロードスター|M6|ファミリア|ボンゴ|スクラム|トリビュート|ベリーサ|RX-7', v)})
blank_join = int((df[AJ] == '').any(axis=1).sum())
t13 = pd.DataFrame({'項目': ['入庫店舗が「0」', '担当者が未設定', '担当者「サービス」（サ-ビス表記を統合）', '担当職が空欄', '車種「その他(…)」の件数', '加入状況が空欄', '誕生月が空欄', '年齢（生年）の列', '年収の列', '最終来店日・点検履歴の列'],
  '件数・状況': [int((df['入庫店舗'] == '0').sum()), int((df['担当者'] == '(未設定)').sum()), int((df['担当者'] == 'サービス').sum()), int((df['担当職'] == '').sum()),
              int(len(other)), blank_join, int((df['誕生月'] == '').sum()), 'なし（誕生月のみ）', 'なし', 'なし'],
  '対応': ['入庫店舗を修正', '担当者を割り当て', '営業担当の割り当てを検討', '任意項目のため対応不要', '車名を統一。マツダ車は正しい車種に修正', '加入状況を入力', '誕生月を入力（11_誕生月の合計に含まれない）', '生年を取得すると年代別配信が可能', 'アンケートで「ご予算・月々のお支払い希望」を取得', '休眠判定に必要。基幹システムから連携を検討']})
LCOLS = ['No', 'お名前(姓)', 'LINEユーザ名', '店舗', '担当者', '車種', '車格', '車種区分', '提案車種', '登録日', '車検日', '車検まで(月)', '次回車検の回数', '任意保険の加入', 'パックdeメンテの加入', 'スカイプランの加入', '優先度', '暫定価格帯セグメント(仮説)', '次に送るステップ', '1回目配信日', '2回目配信日', '3回目配信日', '4回目配信日', '複数台保有']
lst = df[LCOLS].sort_values(['優先度', '車検日']).reset_index(drop=True)

# ---------- 改善点（数値から生成） ----------
def pct(x): return f'{x:.1%}'
s6df = df[s6]
cnt = lambda q: int(q.sum())
disc = t2[t2['車種区分'].isin(['販売終了', '旧名称'])].head(2)
low_pack = t2[(t2['件数'] >= 20)].sort_values('パックdeメンテ加入率').head(3)
mini = t3[t3['車格'] == 'ミニバン'].iloc[0] if (t3['車格'] == 'ミニバン').any() else None
comp = t3[t3['車格'] == 'コンパクト'].iloc[0]
u2 = df[df['車検時期'] == '1) 2か月以内']
over10 = df[df['次回車検の回数'] == '10年超']
first = int((df['次回車検の回数'] == '初回車検(3年)').sum())
peak = t6.sum(axis=1).idxmax(); peakn = int(t6.sum(axis=1).max())
nonpeak = t6.sum(axis=1).iloc[1:-1].nsmallest(2)
aA = df[df['優先度'].str.startswith('A')]
aA_store = aA['店舗'].value_counts()
ast = df[df['優先度'].str.startswith('A')].groupby(['店舗', '担当者']).size()
ast_real = ast[[k[1] not in ('サービス', '(未設定)') for k in ast.index]]
maxs, mins = ast_real.sort_values(ascending=False), ast_real[ast_real.index.map(lambda k: staff.set_index(['店舗', '担当者']).loc[k, '件数'] >= 100)].sort_values()
svc = df[(df['担当者'] == 'サービス')]; svcA = int(svc['優先度'].str.startswith('A').sum())
a1s = df[df['優先度'].str.startswith('A1')].groupby('担当者').size().sort_values(ascending=False)
reg_y = pd.to_datetime(df['お友だち登録日']).dt.year
old = reg_y.value_counts().sort_index()
first2 = old.index[:2]; old_n = int(old.iloc[:2].sum())
low_year = old.iloc[2:-1].idxmin(); low_year_n = int(old.loc[low_year])
IMP = {
 '01_概要': [f"任意保険（店舗扱い）の加入率が{pct(R['ins_rate'])}です。車検入庫時に保険の見直しを提案します。",
           f"スカイプランの加入率は{pct(R['sky_rate'])}です。新車商談で残価設定型プランの提案を標準の流れに組み込みます。",
           '年齢・年収・最終来店日のデータがないため、項目の追加を検討します（13_データ品質を参照）。'],
 '02_車種別': ['販売終了・旧名称の車種（' + '、'.join(f"{r['車種']}{r['件数']}件" for _, r in disc.iterrows()) + 'など）向けに、後継車種を案内する車種別メッセージを用意します。',
           'パックdeメンテの加入率が低い車種（' + '、'.join(f"{r['車種']} {r['パックdeメンテ加入率']:.0%}" for _, r in low_pack.iterrows()) + '）に、点検時の加入案内を行います（20件以上の車種で比較）。',
           f"「その他(…)」の{len(other)}件は表記を統一し、マツダ車（{'、'.join(mazda_in_other[:3])}など）は正しい車種に修正します。"],
 '03_車格別': ([f"ミニバン{mini['件数']}件は平均{mini['平均経過年数']}年と古く、マツダにはミニバンの後継車がないため、他メーカーに流れやすい層です（仮説）。CX-80・CX-5の試乗案内を優先します。"] if mini is not None else []) +
          [f"コンパクト{comp['件数']}件のうち{comp['販売終了旧名称']}件が旧名称（デミオ）です。MAZDA2への乗り換え案内の対象にします。"],
 '04_店舗×車検時期': [f"車検まで2か月以内の{len(u2)}件（" + '、'.join(f'{s}{c}件' for s, c in u2['店舗'].value_counts().items()) + '）は、車検予約が済んでいなければ至急連絡します。',
           f"{STORES[0]}の件数は{STORES[1]}の約{R['stores'][STORES[0]] / R['stores'][STORES[1]]:.1f}倍です。入庫枠と担当者の負荷を店舗ごとに確認します。"],
 '05_次回車検の回数×店舗': [f"10年超の車が{len(over10)}件（{len(over10) / n:.1%}、うち{over10['店舗'].value_counts().index[0]}{over10['店舗'].value_counts().iloc[0]}件）あります。乗り換え提案と、乗り続ける方向けの点検メニューの両方を用意します。",
           f'初回車検の{first}件には、パックdeメンテの継続・延長を案内します。'],
 '06_車検カレンダー12か月': [f"{peak[:4]}年{int(peak[5:])}月に{peakn}件が集中しています。2〜3か月前から試乗枠、1〜2か月前から車検の入庫枠を前もって確保します。",
           '件数の少ない月（' + '・'.join(f"{int(k[5:])}月{v}件" for k, v in nonpeak.items()) + '）に向けて早めの入庫を案内し、作業量をならします（早期受検の条件は制度を確認のうえ案内）。'],
 '07_優先度×店舗': [f"優先度Aは{len(aA)}件（全体の{len(aA) / n:.1%}、うち{aA_store.index[0]}{aA_store.iloc[0]}件）です。担当者ごとに週単位の連絡件数の目標を決めて進めます。",
           f"C2（車検13か月以上）の{cnt(df['優先度'].str.startswith('C2'))}件は、配信を月1〜2回に抑えてブロックを防ぎます。"],
 '08_担当者別': [f"優先度Aの件数に担当者間で差があります（{maxs.index[0][1]}{maxs.iloc[0]}件・{maxs.index[1][1]}{maxs.iloc[1]}件に対し、{mins.index[0][1]}{mins.iloc[0]}件）。フォローの分担を見直します（担当100件以上の方で比較）。",
           f"担当者「サービス」（優先度A {svcA}件）と(未設定)の{cnt(df['担当者'] == '(未設定)')}件に、営業担当を割り当てます。",
           f"{a1s.index[0]}さんはA1（スカイプラン満了期）が{a1s.iloc[0]}件と多いため、満了時の相談を優先して進めます。"],
 '09_加入状況×車検時期': [f"車検6か月以内で任意保険「いいえ」が{cnt(s6df['任意保険の加入'] == 'いいえ')}件、パックdeメンテ「いいえ」が{cnt(s6df['パックdeメンテの加入'] == 'いいえ')}件あります。入庫時の案内対象にします。"] +
          ([f'加入状況が空欄の{blank_join}件を修正します。'] if blank_join else []),
 '10_暫定価格帯セグメント': ['年収データがなく、保有車から想定した暫定（仮説）のセグメントです。LINEアンケートで「ご予算・月々のお支払い希望」を集めて置き換えます。',
           '601万〜700万円には、ミニバン・セダン・スポーツなど用途の違う車種が含まれます。家族の人数や用途で分けて配信します。'],
 '11_誕生月': ['生年（年齢）のデータがないため、年代別の配信ができません。生年の取得を検討します。', 'お誕生日メッセージの特典の有無・内容を、店舗間で統一します。'],
 '12_友だち登録年': [f"{first2[0]}〜{first2[1]}年に登録した方が{old_n}件（{old_n / n:.1%}）です。ブロック率を確認し、配信の頻度を調整します。",
           f"{low_year}年の登録は{low_year_n}件と少ないため、納車時・点検時のお友だち登録の案内を強化します。"],
 '13_データ品質': ['表の「対応」列のとおりです。すぐに対応するもの：入庫店舗「0」と担当者未設定の修正。次に対応するもの：車種の表記統一、生年・年収・最終来店日の項目追加。'],
 '20_配信リスト': [f"複数台保有の方（{R['multi']}人）は、同じ日に2通届かないよう、配信前に重複を除きます。", '配信結果（クリック・予約・成約）を記録する列を追加すると、配信ごとの効果を測れます。'],
}

# ---------- 4. Excel ----------
FONT = 'BIZ UDPゴシック'
GREEN, LIGHT, YEL, HL = (PatternFill('solid', fgColor=c) for c in ('FF00CC66', 'FFE2F7EC', 'FFFFF2CC', 'FFFFFF00'))
THIN, DBL = Side(style='thin'), Side(style='double')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR = 7; FIRST = 8
LN = len(lst); LR = (FIRST, FIRST + LN - 1)
LIST = '20_配信リスト'
def rng(col): return f"'{LIST}'!${col}${LR[0]}:${col}${LR[1]}"
CI = {c: L(i + 1) for i, c in enumerate(LCOLS)}

wb = openpyxl.Workbook(); wb.remove(wb.active)
def write_table(ws, header, rows, widths, total=None, wrap_cols=(), left_cols=(), fmt=None):
    for j, h in enumerate(header, 1): ws.cell(HDR, j, h)
    for i, row in enumerate(rows):
        for j, v in enumerate(row, 1): ws.cell(FIRST + i, j, v)
    last = FIRST + len(rows) - 1
    if total:
        last += 1
        for j, v in enumerate(total(last), 1):
            if v is not None: ws.cell(last, j, v)
    nh = max(str(h).count('\n') + 1 for h in header)
    ws.row_dimensions[HDR].height = max(30, 15.5 * nh + 6)
    for r in range(HDR, last + 1):
        for j in range(1, len(header) + 1):
            c = ws.cell(r, j); head = r == HDR
            c.font = Font(name=FONT, size=11, bold=head or (total is not None and r == last and j == 1))
            if head: c.fill = GREEN
            c.border = BOX
            c.alignment = Alignment(horizontal='left' if (j in left_cols and not head) else 'center', vertical='center', wrap_text=head or j in wrap_cols)
            if fmt and j in fmt and not head: c.number_format = fmt[j]
        if r > HDR: ws.row_dimensions[r].height = 22.05
    for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w
    return last

def finish(ws, ncol, last):
    title, what, how, items = D[ws.title]
    end_col = ncol
    if items:
        ic = ncol + 2; ws.column_dimensions[L(ncol + 1)].width = 3
        ws.cell(HDR, ic, '項目'); ws.cell(HDR, ic + 1, '項目の説明')
        for i, (k, v) in enumerate(items, 1): ws.cell(HDR + i, ic, k); ws.cell(HDR + i, ic + 1, v)
        for r in range(HDR, HDR + len(items) + 1):
            for c in (ic, ic + 1):
                cell = ws.cell(r, c); head = r == HDR
                cell.font = Font(name=FONT, size=11, bold=head); cell.border = BOX
                if head: cell.fill = GREEN
                cell.alignment = Alignment(horizontal='center' if head else 'left', vertical='center', wrap_text=True)
            if ws.row_dimensions[r].height is None: ws.row_dimensions[r].height = 22.05
        ws.column_dimensions[L(ic)].width = 26; ws.column_dimensions[L(ic + 1)].width = 52
        end_col = ic + 1
    while sum(ws.column_dimensions[L(c)].width or 8.9 for c in range(1, end_col + 1)) < 90: end_col += 1
    width = sum(ws.column_dimensions[L(c)].width or 8.9 for c in range(1, end_col + 1))
    ws['A1'] = f'{ws.title}　{title}'; ws['A1'].font = Font(name=FONT, size=14, bold=True); ws.row_dimensions[1].height = 27
    imp = IMP.get(ws.title) or ['なし']
    imp_txt = imp[0] if len(imp) == 1 else '\n'.join(f'{chr(0x2460 + i)} {t}' for i, t in enumerate(imp))
    for r, (lab, txt, fill) in zip((2, 3, 4), (('分析内容', what, LIGHT), ('見方・活用', how, LIGHT), ('改善点', imp_txt, YEL))):
        ws.cell(r, 1, f'【{lab}】' + ('\n' if '\n' in txt else '') + txt)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=end_col)
        c = ws.cell(r, 1); c.font = Font(name=FONT, size=11, bold=(lab == '改善点')); c.fill = fill
        c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        lines = sum(max(1, -(-len(p) * 2 // int(width * 0.95))) for p in c.value.split('\n'))
        ws.row_dimensions[r].height = max(22.05, 16.5 * lines + 6)
    ws['A5'] = f'基準日：{BASE:%Y年%m月%d日}（CSV出力日）　元データ：{SRC}' + ('　※表の右側に項目の説明があります' if items else '')
    ws['A5'].font = Font(name=FONT, size=9, color='FF666666'); ws.row_dimensions[5].height = 18; ws.row_dimensions[6].height = 8
    ws.freeze_panes = f'C{FIRST}' if ws.title == LIST else f'A{FIRST}'

def sum_total(cols, label='合計', first_col_label=True):
    def f(last):
        row = ['合計'] + [None] * (cols - 1)
        return row
    return f

# 01
ws = wb.create_sheet('01_概要')
abs_n = f'$B${FIRST}'
rows01 = [
 ('お友だち（車両）件数', f'=COUNTA({rng("A")})', '0', 'CSVの行数。1件＝お友だち×車両1台'),
 ('人数（電話番号で名寄せ）', R['persons'], '0', '同じ電話番号を1人として数えた人数（配信リストに電話番号がないため値で記載）'),
 ('複数台保有の人数', R['multi'], '0', '2台以上登録がある人数。同じ日に2通届かないよう配信時に重複を除く'),
] + [(s, f'=COUNTIF({rng(CI["店舗"])},"{s}")', '0', f'担当店舗（管理者店舗）が{s}の件数') for s in STORES] + [
 ('任意保険 加入率', f'=COUNTIF({rng(CI["任意保険の加入"])},"はい")/{abs_n}', '0.0%', '任意保険の加入が「はい」の割合（店舗扱いの保険）'),
 ('パックdeメンテ 加入率', f'=COUNTIF({rng(CI["パックdeメンテの加入"])},"はい")/{abs_n}', '0.0%', 'パックdeメンテの加入が「はい」の割合'),
 ('スカイプラン 加入率', f'=COUNTIF({rng(CI["スカイプランの加入"])},"はい")/{abs_n}', '0.0%', '残価設定型クレジット「スカイプラン」の加入が「はい」の割合'),
 ('平均経過年数（現在）', round(R['avg_age'], 1), '0.0', f'車の登録日から基準日（{BASE:%Y/%m/%d}）までの年数の平均（値で記載）'),
 ('車検6か月以内（件）', f'=COUNTIF({rng(CI["車検まで(月)"])},"<=6")', '0', '車検日まで6か月以内の件数。優先配信・ステップ配信の対象'),
 ('車検12か月以内（件）', f'=COUNTIF({rng(CI["車検まで(月)"])},"<=12")', '0', '車検日まで12か月以内の件数'),
]
last = write_table(ws, ['項目', '値', '説明'], [(a, b, d) for a, b, _, d in rows01], [25, 15.8, 66], wrap_cols=(3,), left_cols=(3,))
for i, (_, _, f, _) in enumerate(rows01): ws.cell(FIRST + i, 2).number_format = f
finish(ws, 3, last)

# 02
ws = wb.create_sheet('02_車種別')
k = len(t2); tot = FIRST + k
rows = [(r['車種'], r['件数'], f'=B{FIRST + i}/$B${tot}', r['平均経過年数'], r['車検6か月以内'], r['スカイプラン加入'], r['パックdeメンテ加入率'], r['車格'], r['車種区分'], r['提案車種']) for i, (_, r) in enumerate(t2.iterrows())]
last = write_table(ws, ['車種', '件数', '構成比', '平均経過年数', '車検\n6か月以内', 'スカイプラン加入', 'パックdeメンテ\n加入率', '車格', '車種区分', '提案車種'], rows,
                   [32.4] + [15.8] * 8 + [28.8], total=lambda r: ['合計', f'=SUM(B{FIRST}:B{r - 1})', f'=SUM(C{FIRST}:C{r - 1})', None, f'=SUM(E{FIRST}:E{r - 1})', f'=SUM(F{FIRST}:F{r - 1})'], fmt={3: '0.00%', 7: '0%'})
for r in range(FIRST, tot):
    if ws.cell(r, 9).value == '販売終了': ws.cell(r, 9).fill = HL
finish(ws, 10, last)

# 03
ws = wb.create_sheet('03_車格別')
k = len(t3); tot = FIRST + k
rows = [(r['車格'], r['件数'], f'=B{FIRST + i}/$B${tot}', r['平均経過年数'], r['車検6か月以内'], r['販売終了旧名称']) for i, (_, r) in enumerate(t3.iterrows())]
last = write_table(ws, ['車格', '件数', '構成比\n（全体）', '平均経過年数', '車検\n6か月以内', '販売終了\n旧名称'], rows, [25.8] + [15.8] * 5,
                   total=lambda r: ['合計', f'=SUM(B{FIRST}:B{r - 1})', f'=SUM(C{FIRST}:C{r - 1})', None, f'=SUM(E{FIRST}:E{r - 1})', f'=SUM(F{FIRST}:F{r - 1})'], fmt={3: '0.00%'})
finish(ws, 6, last)

def cross(ws, name0, t, w0, headers=None, double_total=False):
    cols = list(t.columns); nc = len(cols) + 2
    rows = [[idx] + [int(v) for v in t.loc[idx]] + [f'=SUM(B{FIRST + i}:{L(nc - 1)}{FIRST + i})'] for i, idx in enumerate(t.index)]
    hdr = [name0] + (headers or [str(c) for c in cols]) + ['合計']
    last = write_table(ws, hdr, rows, [w0] + [15.8] * (nc - 1), total=lambda r: ['合計'] + [f'=SUM({L(j)}{FIRST}:{L(j)}{r - 1})' for j in range(2, nc + 1)])
    if double_total:
        for r in range(HDR, last + 1):
            c = ws.cell(r, nc); c.border = Border(left=DBL, right=THIN, top=THIN, bottom=THIN)
    finish(ws, nc, last)

cross(wb.create_sheet('04_店舗×車検時期'), '店舗', t4, 15.8, headers=[b.replace(' ', ' \n') for b in BUCK])
cross(wb.create_sheet('05_次回車検の回数×店舗'), '次回車検の回数', t5, 15.8)
cross(wb.create_sheet('06_車検カレンダー12か月'), '満了月', t6, 15.8)
cross(wb.create_sheet('07_優先度×店舗'), '優先度', t7, 40.8)
for r in range(FIRST, FIRST + len(PR)): wb['07_優先度×店舗'].cell(r, 1).alignment = Alignment(horizontal='left', vertical='center')

# 08
ws = wb.create_sheet('08_担当者別')
PAT = ['A1*', 'A2*', 'A3*', 'B*', 'C1*', 'C2*']
rows = []
for i, (_, r) in enumerate(staff.iterrows()):
    x = FIRST + i
    base = f'{rng(CI["店舗"])},$A{x},{rng(CI["担当者"])},$B{x}'
    rows.append([r['店舗'], r['担当者'], f'=COUNTIFS({base})', f'=COUNTIFS({base},{rng(CI["車検まで(月)"])},"<=6")', f'=SUM(F{x}:H{x})'] +
                [f'=COUNTIFS({base},{rng(CI["優先度"])},"{p}")' for p in PAT] + [f'=IF(SUM(F{x}:K{x})=C{x},"OK","要確認")'])
last = write_table(ws, ['店舗', '担当者', '件数', '車検\n6か月以内', '優先度A\n計(A1〜A3)', 'A1\nスカイプラン\n満了期', 'A2\n販売終了\n車種', 'A3\n7年以上', 'B\n車検予約', 'C1\n準備層\n7〜12か月', 'C2\n関係維持\n13か月以上', '件数との\n一致確認'],
                   rows, [15.8] * 12, total=lambda r: ['合計', None] + [f'=SUM({L(j)}{FIRST}:{L(j)}{r - 1})' for j in range(3, 12)] + [f'=IF(SUM(F{r}:K{r})=C{r},"OK","要確認")'])
ws.merge_cells(f'A{last}:B{last}')
finish(ws, 12, last)

# 09
ws = wb.create_sheet('09_加入状況×車検時期')
cols9 = list(t9.columns)
rows = [[b] + [int(t9.loc[b, c]) for c in cols9] for b in BUCK]
hdr = ['車検時期'] + [f'{a}:\n{b or "（空欄）"}' for a, b in cols9]
last = write_table(ws, hdr, rows, [22.8] + [15.8] * len(cols9), total=lambda r: ['合計'] + [f'=SUM({L(j)}{FIRST}:{L(j)}{r - 1})' for j in range(2, len(cols9) + 2)])
finish(ws, len(cols9) + 1, last)

# 10
ws = wb.create_sheet('10_暫定価格帯セグメント')
rows = [[s, seg_models.get(s, '')] + [int(v) for v in t10.loc[s]] + [f'=SUM(C{FIRST + i}:G{FIRST + i})'] for i, s in enumerate(SEGS)]
last = write_table(ws, ['暫定価格帯セグメント(仮説)', '対象車種（保有車）'] + BUCK + ['合計'], rows, [28.3, 62] + [15.8] * 6, wrap_cols=(2,), left_cols=(2,),
                   total=lambda r: ['合計', None] + [f'=SUM({L(j)}{FIRST}:{L(j)}{r - 1})' for j in range(3, 9)])
for i in range(len(SEGS)):
    ln = len(str(ws.cell(FIRST + i, 2).value)); ws.row_dimensions[FIRST + i].height = max(22.05, 16.5 * -(-ln // 34) + 6)
finish(ws, 8, last)

cross(wb.create_sheet('11_誕生月'), '誕生月', t11, 15.8, double_total=True)
wb['11_誕生月'].cell(HDR, len(STORES) + 2, '対象者\n合計')
t12.index.name = 'お友だち登録年'
cross(wb.create_sheet('12_友だち登録年'), 'お友だち登録年', t12, 15.8, double_total=True)
wb['12_友だち登録年'].cell(HDR, len(STORES) + 2, '登録数\n合計')

# 13
ws = wb.create_sheet('13_データ品質')
last = write_table(ws, ['項目', '件数・状況', '対応'], t13.values.tolist(), [38.8, 20.8, 50.8], left_cols=(1, 3), wrap_cols=(3,))
finish(ws, 3, last)

# 20
ws = wb.create_sheet(LIST)
vals = lst.values.tolist()
widths = [14.4, 21.6, 40, 8, 9, 32.4, 18, 8, 28.8, 18, 12.6, 12.6, 14.4, 12.6, 19.8, 16.2, 41.8, 30.8, 48.8, 15.8, 15.8, 15.8, 15.8, 12]
for j, h in enumerate(LCOLS, 1): ws.cell(HDR, j, h)
for i, row in enumerate(vals):
    for j, v in enumerate(row, 1):
        c = ws.cell(FIRST + i, j, float(v) if j == 12 else v)
        c.font = Font(name=FONT, size=11); c.border = BOX; c.alignment = Alignment(horizontal='center', vertical='center')
for j in range(1, len(LCOLS) + 1):
    c = ws.cell(HDR, j); c.font = Font(name=FONT, size=11, bold=True); c.fill = GREEN; c.border = BOX
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.column_dimensions[L(j)].width = widths[j - 1]
ws.row_dimensions[HDR].height = 30
ws.auto_filter.ref = f'A{HDR}:{L(len(LCOLS))}{LR[1]}'
finish(ws, len(LCOLS), LR[1])

wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(out_xlsx)

# ---------- 計画書・メッセージ用の数値 ----------
out = dict(R, base=str(BASE.date()), src=SRC, unclassified=unclassified,
  prio={p: {s: int(t7.loc[p, s]) for s in STORES} for p in PR},
  urgent=int(df['至急'].sum()), urgent_store=df[df['至急']]['店舗'].value_counts().to_dict(), urgent_by=df[df['至急']]['優先度'].str[:2].str.strip().value_counts().to_dict(),
  feat={p[:2].strip(): {'models': df[df['優先度'] == p]['車種'].value_counts().head(5).to_dict(), 'kaisu': df[df['優先度'] == p]['次回車検の回数'].value_counts().to_dict()} for p in PR},
  x1=cnt(s6df['パックdeメンテの加入'] == 'いいえ'), x1_store=s6df[s6df['パックdeメンテの加入'] == 'いいえ']['店舗'].value_counts().to_dict(), x2=cnt(s6df['任意保険の加入'] == 'いいえ'), x2_store=s6df[s6df['任意保険の加入'] == 'いいえ']['店舗'].value_counts().to_dict(),
  bday={m_: int((df['誕生月'] == m_).sum()) for m_ in [f'{i}月' for i in range(1, 13)]},
  cal={k: int(v) for k, v in t6.sum(axis=1).items()}, cal_store={k: {s: int(t6.loc[k, s]) for s in STORES} for k in t6.index},
  steps={n_: pd.to_datetime(df[f'{n_}配信日'], errors='coerce', format='%Y-%m-%d').dt.strftime('%Y-%m').value_counts().sort_index().to_dict() for n_, _ in STEPS},
  nextstep=df['次に送るステップ'].value_counts().to_dict(),
  kakaku=t10.sum(axis=1).to_dict(), kakaku6={s: int(t10.loc[s, BUCK[:3]].sum()) for s in SEGS},
  t2=t2.head(10).to_dict('records'), t3=t3.to_dict('records'), t4={s: {b: int(t4.loc[s, b]) for b in BUCK} for s in STORES},
  t5={o: {s: int(t5.loc[o, s]) for s in STORES} for o in ORD}, imp=IMP, a1_staff=a1s.head(3).to_dict(),
  aA_staff={f'{k[1]}（{k[0]}）': int(v) for k, v in ast.sort_values(ascending=False).items()},
  reg_year={int(k): int(v) for k, v in old.items()}, quality=t13.astype(str).values.tolist())
json.dump(out, open(out_json, 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)
print('OK', out_xlsx, 'unclassified:', unclassified)
