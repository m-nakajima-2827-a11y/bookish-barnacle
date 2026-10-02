"""原稿JSONから、白基調・ミニマルな提案書（.pptx）を生成する。

使い方:
    python build_deck.py <原稿.json> <出力.pptx>

原稿の構造は examples/sample-deck.json を参照。文字列中の **...** は太字になる。
"""
import json
import re
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

INK = RGBColor(0x11, 0x11, 0x11)
SUB = RGBColor(0x8A, 0x8A, 0x8A)
RULE = RGBColor(0xE6, 0xE6, 0xE6)
ACCENT = RGBColor(0x1F, 0x4F, 0xFF)
MUTED_BAR = RGBColor(0xD0, 0xD0, 0xD0)
FONT = "Noto Sans JP"

W, H = 13.333, 7.5
ML, MR, MT, MB = 0.9, 0.9, 0.7, 0.6
CW = W - ML - MR


def _runs(paragraph, text, size, color=INK, bold=False):
    """**...** を太字の run に分けて追加する。"""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        r = paragraph.add_run()
        r.text = part
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold or i % 2 == 1


def text(slide, x, y, w, h, lines, size=14, color=INK, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.4, after=6):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        _runs(p, line, size, color, bold)
    return box


def rule(slide, x, y, w):
    ln = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = RULE
    ln.line.width = Pt(0.75)


def footer(slide, meta, page):
    text(slide, ML, H - MB - 0.1, 6, 0.3, meta.get("client", ""), 9, SUB)
    text(slide, W - MR - 1, H - MB - 0.1, 1, 0.3, f"{page:02d}", 9, SUB,
         align=PP_ALIGN.RIGHT)


def section(slide, x, y, w, label, items, numbered=False):
    text(slide, x, y, w, 0.3, label, 10, SUB)
    lines = [f"{i}. {s}" if numbered else f"・{s}"
             for i, s in enumerate(items[:2], 1)]
    text(slide, x, y + 0.35, w, 1.3, lines, 14)


def fit_size(s, w, max_pt=60):
    """1行に収まるフォントサイズ（全角1em、半角0.6emで概算）。"""
    em = sum(0.6 if ord(ch) < 0x2E80 else 1.0 for ch in s) or 1
    return max(28, min(max_pt, int(w * 72 / em * 0.92)))


def metric(slide, x, y, w, m):
    text(slide, x, y, w, 1.1, m["value"], fit_size(m["value"], w), ACCENT,
         bold=True, anchor=MSO_ANCHOR.BOTTOM, spacing=1.0)
    text(slide, x, y + 1.15, w, 0.4, m["label"], 12, INK)
    if m.get("note"):
        text(slide, x, y + 1.55, w, 0.6, m["note"], 9, SUB)


def chart(slide, x, y, w, h, c):
    data = CategoryChartData()
    data.categories = c["categories"]
    data.add_series("", c["values"])
    kind = XL_CHART_TYPE.LINE_MARKERS if c.get("type") == "line" \
        else XL_CHART_TYPE.COLUMN_CLUSTERED
    if c.get("title"):
        text(slide, x, y, w, 0.3, c["title"], 10, SUB)
    gf = slide.shapes.add_chart(
        kind, Inches(x), Inches(y + 0.35), Inches(w), Inches(h - 0.8), data)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = FONT
    ch.font.size = Pt(10)
    ch.font.color.rgb = SUB
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.format.line.color.rgb = RULE
    ca.has_major_gridlines = False
    plot = ch.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0.0"' + c.get("unit", "") + '"' \
        if any(isinstance(v, float) and not v.is_integer() for v in c["values"]) \
        else '0"' + c.get("unit", "") + '"'
    dl.number_format_is_linked = False
    dl.font.size = Pt(11)
    dl.font.color.rgb = INK
    hi = c.get("highlight", len(c["values"]) - 1)
    series = plot.series[0]
    if kind == XL_CHART_TYPE.COLUMN_CLUSTERED:
        plot.gap_width = 80
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        for i in range(len(c["values"])):
            pt = series.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = ACCENT if i == hi else MUTED_BAR
    else:
        series.format.line.color.rgb = ACCENT
        series.format.line.width = Pt(2)
        dl.position = XL_LABEL_POSITION.ABOVE
    if c.get("note"):
        text(slide, x, y + h - 0.4, w, 0.3, c["note"], 9, SUB)


def slide_cover(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    m = d["meta"]
    text(s, ML, 2.6, CW, 1.0, m["title"], 44, bold=True)
    if m.get("subtitle"):
        text(s, ML, 3.75, CW, 0.5, m["subtitle"], 16, SUB)
    rule(s, ML, 4.5, 1.2)
    text(s, ML, 4.75, CW, 1.2,
         [m.get("client", ""), f'{m.get("date", "")}　{m.get("author", "")}'],
         12, SUB, spacing=1.2, after=4)


def slide_overview(prs, d, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    o = d["overview"]
    text(s, ML, MT, CW, 0.3, "全体サマリー", 10, SUB)
    text(s, ML, MT + 0.35, CW, 0.7, o["title"], 28, bold=True)
    text(s, ML, MT + 1.1, CW, 0.5, o["core"], 16, SUB)
    top = MT + 2.0
    cols = [("No", 0.6), ("提案", 4.6), ("優先度", 1.1), ("期間", 1.5), ("KPI", CW - 7.8)]
    x = ML
    for name, w in cols:
        text(s, x, top, w, 0.3, name, 10, SUB)
        x += w
    rule(s, ML, top + 0.4, CW)
    rows = o["proposals"][:5]
    row_h = min(0.75, (H - MB - 0.5 - top - 0.5) / max(len(rows), 1))
    for r, p in enumerate(rows):
        y = top + 0.55 + r * row_h
        vals = [f'{r + 1:02d}', p["title"], p["priority"], p["period"], p["kpi"]]
        x = ML
        for (name, w), v in zip(cols, vals):
            text(s, x, y, w - 0.15, row_h, v, 14,
                 ACCENT if name == "No" else INK, bold=name in ("No", "提案"))
            x += w
        rule(s, ML, y + row_h - 0.12, CW)
    footer(s, d["meta"], page)


def slide_proposal(prs, d, p, no, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, ML, MT, 3, 0.3, f"提案 {no:02d}", 10, ACCENT)
    text(s, ML, MT + 0.35, CW, 0.7, p["title"], 28, bold=True)
    text(s, ML, MT + 1.1, CW, 0.5, p["core"], 16, SUB)
    rule(s, ML, MT + 1.75, CW)

    gx, gy, gw = ML, MT + 2.05, 8.3
    cw = (gw - 0.4) / 2
    section(s, gx, gy, cw, "課題 / WHY", p["why"])
    section(s, gx + cw + 0.4, gy, cw, "解決策 / WHAT", p["what"])
    section(s, gx, gy + 2.0, cw, "効果 / ROI", p["roi"])
    section(s, gx + cw + 0.4, gy + 2.0, cw, "動き / NEXT", p["next"], numbered=True)

    rx = ML + gw + 0.4
    rw = W - MR - rx
    vline = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(rx - 0.2), Inches(gy),
                                   Inches(rx - 0.2), Inches(gy + 3.7))
    vline.line.color.rgb = RULE
    vline.line.width = Pt(0.75)
    if p.get("chart"):
        chart(s, rx + 0.1, gy, rw - 0.1, 3.8, p["chart"])
    elif p.get("metric"):
        metric(s, rx + 0.1, gy + 0.6, rw - 0.1, p["metric"])
    footer(s, d["meta"], page)


def slide_next(prs, d, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, ML, MT, CW, 0.3, "次の動き / NEXT ACTION", 10, SUB)
    text(s, ML, MT + 0.35, CW, 0.7, "直近の3ステップ" if len(d["next_actions"]) == 3
         else "直近のステップ", 28, bold=True)
    top = MT + 1.7
    rows = d["next_actions"][:5]
    for i, a in enumerate(rows):
        y = top + i * 0.85
        text(s, ML, y, 0.8, 0.6, f"{i + 1}.", 24, ACCENT, bold=True)
        text(s, ML + 0.8, y + 0.08, 6.8, 0.6, a["step"], 16, bold=True)
        text(s, ML + 7.8, y + 0.12, 2.0, 0.5, a.get("owner", ""), 12, SUB)
        text(s, ML + 9.9, y + 0.12, CW - 9.9, 0.5, a.get("due", ""), 12, SUB)
        rule(s, ML, y + 0.7, CW)
    footer(s, d["meta"], page)


def slide_appendix(prs, d, ap, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, ML, MT, CW, 0.3, "付録 / APPENDIX", 10, SUB)
    text(s, ML, MT + 0.35, CW, 0.7, ap["title"], 28, bold=True)
    head, *rows = ap["rows"]
    widths = ap.get("widths") or [CW / len(head)] * len(head)
    top = MT + 1.5
    x = ML
    for h_, w in zip(head, widths):
        text(s, x, top, w - 0.15, 0.3, h_, 10, SUB)
        x += w
    rule(s, ML, top + 0.38, CW)
    row_h = min(0.6, (H - MB - 0.5 - top - 0.5) / max(len(rows), 1))
    for r, row in enumerate(rows):
        y = top + 0.5 + r * row_h
        x = ML
        for v, w in zip(row, widths):
            text(s, x, y, w - 0.15, row_h, v, 11, spacing=1.2, after=0)
            x += w
    if ap.get("note"):
        text(s, ML, H - MB - 0.5, CW, 0.3, ap["note"], 9, SUB)
    footer(s, d["meta"], page)


def build(d, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    slide_cover(prs, d)
    page = 2
    slide_overview(prs, d, page)
    for no, p in enumerate(d["proposals"][:5], 1):
        page += 1
        slide_proposal(prs, d, p, no, page)
    if d.get("next_actions"):
        page += 1
        slide_next(prs, d, page)
    for ap in d.get("appendix", []):
        page += 1
        slide_appendix(prs, d, ap, page)
    prs.save(out)
    return page


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: python build_deck.py <deck.json> <out.pptx>")
    with open(sys.argv[1], encoding="utf-8") as f:
        deck = json.load(f)
    n = len(deck.get("proposals", []))
    if not 3 <= n <= 5:
        sys.exit(f"提案は3〜5点にしてください（現在 {n} 点）")
    pages = build(deck, sys.argv[2])
    print(f"saved {sys.argv[2]} ({pages} slides)")
